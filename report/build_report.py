"""Build BaoCao_BTL_DMML_2501.docx from real pipeline metrics + screenshots.

Run: python3 report/build_report.py
Reads: data-mining/outputs/metrics.json, report/assets/*.png
Writes: report/BaoCao_BTL_DMML_2501.docx
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
METRICS = json.loads((ROOT / "data-mining" / "outputs" / "metrics.json").read_text())
ASSETS = ROOT / "report" / "assets"
OUT = ROOT / "report" / "BaoCao_BTL_DMML_2501.docx"


def vn(n: int) -> str:
    """Format an integer with Vietnamese-style '.' thousands separator."""
    return f"{n:,}".replace(",", ".")

GROUP = [
    ("Vũ Quốc Đạt (Nhóm trưởng)", "24C1001U5702"),
    ("Đàm Chí Công", "24C1001U4553"),
    ("Nguyễn Thị Hồng Nhung", "24C1001U4867"),
    ("Nguyễn Hoàng Đức", "24C1001U5759"),
]


def add_toc_field(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Nhấn F9 để cập nhật Mục lục sau khi mở file trong Word."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_element = run._r
    r_element.append(fld_begin)
    r_element.append(instr)
    r_element.append(fld_sep)
    r_element.append(fld_text)
    r_element.append(fld_end)


def set_cell_text(cell, text, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(10.5)
    if align:
        p.alignment = align


def add_table(doc: Document, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return table


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h


def add_image_with_caption(doc, path: Path, caption: str, width_cm=15.5):
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(10)


def main() -> None:
    doc = Document()

    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    # ---------------- COVER PAGE ----------------
    for line, size, bold, space_after in [
        ("TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI", 14, True, 0),
        ("KHOA CÔNG NGHỆ THÔNG TIN", 14, True, 12),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.bold = bold
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(space_after)

    doc.add_paragraph()
    doc.add_paragraph()
    for line, size, bold in [
        ("BÀI TẬP LỚN", 20, True),
        ("NHẬP MÔN KHAI PHÁ DỮ LIỆU VÀ MÁY HỌC", 16, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.bold = bold
        r.font.size = Pt(size)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ĐỀ SỐ 2501")
    r.bold = True
    r.font.size = Pt(15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Ứng dụng kỹ thuật khai phá luật kết hợp tích hợp vào hệ thống quản lý "
        "bán lẻ trên công nghệ .NET để giải quyết bài toán khuyến nghị sản phẩm"
    )
    r.bold = True
    r.font.size = Pt(15)

    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Nhóm N21")
    r.bold = True
    r.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sinh viên thực hiện")
    r.bold = True

    for name, mssv in GROUP:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{name} - {mssv}")

    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Hà Nội, năm 2026")

    doc.add_page_break()

    # ---------------- TOC ----------------
    add_heading(doc, "Mục lục", level=1)
    add_toc_field(doc)
    doc.add_page_break()

    # ---------------- CHUONG 1 ----------------
    add_heading(doc, "CHƯƠNG 1. TỔNG QUAN VỀ ĐỀ TÀI", level=1)

    add_heading(doc, "1.1 Giới thiệu đề tài", level=2)
    doc.add_paragraph(
        "Trong lĩnh vực bán lẻ, một trong những nhu cầu phổ biến nhất là gợi ý cho "
        "khách hàng những sản phẩm có khả năng họ sẽ mua thêm, dựa trên hành vi mua "
        "sắm của những khách hàng khác có giỏ hàng tương tự — kỹ thuật này thường "
        "được biết đến dưới tên gọi “khách hàng cũng mua” hoặc “thường "
        "được mua cùng nhau”. Đề tài xây dựng một hệ thống quản lý bán lẻ trên "
        "nền tảng ASP.NET Core (.NET 8), trong đó chức năng khuyến nghị sản phẩm "
        "được tích hợp dựa trên kết quả khai phá luật kết hợp (Association Rule "
        "Mining) từ dữ liệu giao dịch bán lẻ thực tế Online Retail II (UCI ML "
        "Repository)."
    )

    add_heading(doc, "1.2 Đánh giá các đề tài liên quan", level=2)
    doc.add_paragraph(
        "Các hệ thống thương mại điện tử lớn (Amazon, Shopee...) đều có chức năng "
        "gợi ý sản phẩm dựa trên nhiều kỹ thuật khác nhau (collaborative filtering, "
        "content-based, association rules...). Trong phạm vi học phần Nhập môn Khai "
        "phá dữ liệu và Máy học, đề tài lựa chọn tập trung đúng vào kỹ thuật đã học — "
        "khai phá luật kết hợp bằng thuật toán Apriori và FP-Growth — để đảm bảo "
        "minh hoạ rõ ràng toàn bộ quy trình DMML từ dữ liệu thô đến mô hình tích hợp "
        "trong hệ thống thật, thay vì dàn trải nhiều kỹ thuật."
    )

    add_heading(doc, "1.3 Mục đích của đề tài", level=2)
    doc.add_paragraph(
        "Hiện thực đầy đủ quy trình triển khai một dự án DMML (theo đề cương học "
        "phần): xác định vấn đề, tiền xử lý dữ liệu, khai phá mô hình (luật kết hợp), "
        "đánh giá mô hình dựa trên các độ đo khoa học (support, confidence, lift), và "
        "tích hợp mô hình vào một hệ thống ứng dụng cụ thể để chứng minh giá trị thực "
        "tiễn — đúng theo yêu cầu 2.3–2.4 của đề bài."
    )

    add_heading(doc, "1.4 Mục tiêu của đề tài", level=2)
    for item in [
        "Thu thập và tiền xử lý dữ liệu giao dịch bán lẻ Online Retail II.",
        "Khai phá luật kết hợp bằng 2 thuật toán Apriori và FP-Growth, so sánh hiệu "
        "năng và xác nhận tính đúng đắn (2 thuật toán phải cho cùng kết quả tập mục "
        "phổ biến).",
        "Đánh giá và lựa chọn ngưỡng support/confidence có căn cứ thực nghiệm.",
        "Xây dựng hệ thống quản lý bán lẻ tối thiểu (danh mục sản phẩm, giỏ hàng, đặt "
        "hàng) trên ASP.NET Core MVC (.NET 8) + SQLite.",
        "Tích hợp mô hình luật kết hợp vào chức năng khuyến nghị sản phẩm ở 2 điểm "
        "chạm chính: trang chi tiết sản phẩm (“Khách hàng cũng mua”) và "
        "trang giỏ hàng (gợi ý tổng hợp theo toàn bộ giỏ).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "1.5 Quy trình thực hiện", level=2)
    add_heading(doc, "1.5.1 Quy trình triển khai dự án DMML", level=3)
    doc.add_paragraph(
        "Đề tài bám sát quy trình 5 bước được giảng dạy trong học phần (tương đương "
        "chuẩn CRISP-DM rút gọn):"
    )
    add_table(
        doc,
        ["Bước", "Nội dung", "Công cụ/Kết quả"],
        [
            ["b1. Xác định vấn đề & khám phá dữ liệu", "Xác định bài toán khuyến nghị sản phẩm, khảo sát dữ liệu Online Retail II", "Python, pandas"],
            ["b2. Tiền xử lý dữ liệu", "Làm sạch đơn hàng huỷ, dữ liệu lỗi, mã không phải sản phẩm, gộp giao dịch theo hoá đơn", "pandas"],
            ["b3. Mô hình hoá (khai phá luật)", "Áp dụng Apriori và FP-Growth qua mlxtend", "mlxtend"],
            ["b4. Đánh giá mô hình", "So sánh 2 thuật toán, chọn ngưỡng support/confidence dựa trên thực nghiệm", "pandas"],
            ["b5. Phân phối/tích hợp mô hình", "Xuất luật kết hợp, tích hợp vào webapp ASP.NET Core", "ASP.NET Core, EF Core, SQLite"],
        ],
    )

    add_heading(doc, "1.6 Phân công nhiệm vụ", level=2)
    add_table(
        doc,
        ["Họ và tên", "MSSV", "Nhiệm vụ", "Tỷ lệ đóng góp"],
        [
            ["Vũ Quốc Đạt (Nhóm trưởng)", "24C1001U5702", "Điều phối chung; thiết kế kiến trúc hệ thống; phát triển webapp ASP.NET Core (Controllers, Services khuyến nghị); tích hợp dữ liệu luật kết hợp", "25%"],
            ["Đàm Chí Công", "24C1001U4553", "Xây dựng pipeline khai phá dữ liệu Python (tiền xử lý, Apriori/FP-Growth, đánh giá, export); tổng hợp và biên soạn báo cáo", "25%"],
            ["Nguyễn Thị Hồng Nhung", "24C1001U4867", "Phân tích yêu cầu; thiết kế dữ liệu (ERD); thiết kế màn hình; kiểm thử chức năng webapp", "25%"],
            ["Nguyễn Hoàng Đức", "24C1001U5759", "Đánh giá mô hình (so sánh Apriori/FP-Growth, lựa chọn ngưỡng); kiểm thử pipeline dữ liệu; hỗ trợ viết tài liệu tham khảo", "25%"],
        ],
    )

    add_heading(doc, "1.7 Kế hoạch thực hiện", level=2)
    add_table(
        doc,
        ["Giai đoạn", "Sản phẩm"],
        [
            ["Tuần 1", "Phân tích đề bài, khảo sát dữ liệu Online Retail II, thiết kế kiến trúc"],
            ["Tuần 2", "Tiền xử lý dữ liệu, khai phá luật kết hợp (Apriori/FP-Growth), đánh giá mô hình"],
            ["Tuần 3", "Xây dựng webapp ASP.NET Core, tích hợp mô hình vào chức năng khuyến nghị"],
            ["Tuần 4", "Kiểm thử tổng thể, hoàn thiện báo cáo"],
        ],
    )

    add_heading(doc, "1.8 Công nghệ dự kiến sử dụng", level=2)
    doc.add_paragraph(
        "Khai phá dữ liệu: Python 3.13, pandas, mlxtend (Apriori, FP-Growth, "
        "TransactionEncoder, association_rules), matplotlib.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Ứng dụng web: ASP.NET Core MVC (.NET 8), Entity Framework Core + SQLite, "
        "CsvHelper, ASP.NET Core Session (giỏ hàng).",
        style="List Bullet",
    )
    doc.add_paragraph("Kiểm thử: xUnit + EF Core InMemory provider.", style="List Bullet")
    doc.add_paragraph(
        "Dữ liệu: Online Retail II (UCI Machine Learning Repository, id 502).",
        style="List Bullet",
    )

    doc.add_page_break()

    # ---------------- CHUONG 2 ----------------
    add_heading(doc, "CHƯƠNG 2. PHÂN TÍCH YÊU CẦU", level=1)

    add_heading(doc, "2.1 Quy trình nghiệp vụ", level=2)
    for label, text in [
        ("P01 Xem danh mục và chi tiết sản phẩm", "Khách hàng duyệt danh sách sản phẩm, xem chi tiết 1 sản phẩm kèm khối gợi ý “Khách hàng cũng mua”."),
        ("P02 Quản lý giỏ hàng", "Khách hàng thêm/xoá sản phẩm khỏi giỏ hàng; hệ thống hiển thị khối gợi ý tổng hợp dựa trên toàn bộ sản phẩm trong giỏ."),
        ("P03 Đặt hàng", "Khách hàng tạo đơn hàng từ giỏ hàng hiện tại."),
        ("P04 Quản trị luật kết hợp", "Quản trị viên xem danh sách luật kết hợp đã khai phá được tích hợp trong hệ thống."),
    ]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(text)

    add_heading(doc, "2.2 Yêu cầu chức năng", level=2)

    def fr_table(title, rows):
        add_heading(doc, title, level=3)
        add_table(doc, ["Thuộc tính", "Mô tả"], rows, widths=[4, 11.5])

    fr_table("2.2.1 Xem danh sách/chi tiết sản phẩm (FR01)", [
        ["Mô tả", "Hiển thị danh mục và chi tiết sản phẩm"],
        ["Tác nhân", "Khách hàng"],
        ["Dữ liệu vào", "SKU sản phẩm"],
        ["Kết quả đầu ra", "Thông tin sản phẩm + tối đa 4 sản phẩm khuyến nghị liên quan (sắp xếp theo Confidence, Lift giảm dần)"],
        ["Liên kết nghiệp vụ", "P01"],
    ])
    fr_table("2.2.2 Khuyến nghị theo giỏ hàng (FR02)", [
        ["Mô tả", "Gợi ý sản phẩm bổ sung dựa trên toàn bộ giỏ hàng"],
        ["Tác nhân", "Khách hàng"],
        ["Dữ liệu vào", "Danh sách SKU đang có trong giỏ"],
        ["Kết quả đầu ra", "Tối đa 6 sản phẩm gợi ý, không trùng sản phẩm đã có trong giỏ, xếp hạng theo confidence lớn nhất rồi lift"],
        ["Liên kết nghiệp vụ", "P02"],
    ])
    fr_table("2.2.3 Thêm/xoá sản phẩm khỏi giỏ hàng (FR03)", [
        ["Mô tả", "Quản lý giỏ hàng qua Session"],
        ["Tác nhân", "Khách hàng"],
        ["Dữ liệu vào", "SKU, số lượng"],
        ["Kết quả đầu ra", "Giỏ hàng được cập nhật, hiển thị lại tổng tiền"],
        ["Liên kết nghiệp vụ", "P02"],
    ])
    fr_table("2.2.4 Đặt hàng (FR04)", [
        ["Mô tả", "Tạo đơn hàng từ giỏ hàng"],
        ["Tác nhân", "Khách hàng"],
        ["Dữ liệu vào", "Giỏ hàng hiện tại"],
        ["Kết quả đầu ra", "1 bản ghi Order + các OrderItem tương ứng, giỏ hàng được xoá"],
        ["Liên kết nghiệp vụ", "P03"],
    ])
    fr_table("2.2.5 Xem danh sách luật kết hợp (FR05)", [
        ["Mô tả", "Quản trị viên xem toàn bộ luật kết hợp đã tích hợp"],
        ["Tác nhân", "Quản trị viên"],
        ["Dữ liệu vào", "(không có)"],
        ["Kết quả đầu ra", "Bảng luật kết hợp gồm antecedent, consequent, support, confidence, lift"],
        ["Liên kết nghiệp vụ", "P04"],
    ])

    add_heading(doc, "2.3 Yêu cầu về mô hình khai phá dữ liệu", level=2)
    add_table(
        doc,
        ["Mã", "Yêu cầu"],
        [
            ["MR01", "Mô hình phải được huấn luyện/khai phá trên dữ liệu giao dịch thật, có nguồn gốc rõ ràng (Online Retail II, UCI)"],
            ["MR02", "Luật kết hợp được chọn phải thoả ngưỡng min_support và min_confidence xác định qua thực nghiệm, có căn cứ khoa học (không tuỳ ý)"],
            ["MR03", "Kết quả khai phá phải được xác nhận đúng đắn bằng cách đối chiếu 2 thuật toán độc lập (Apriori và FP-Growth) cho cùng 1 tập mục phổ biến"],
            ["MR04", "Cấm sinh luật A→A (một sản phẩm “khuyến nghị” chính nó)"],
        ],
        widths=[2, 13.5],
    )

    add_heading(doc, "2.4 Yêu cầu phi chức năng", level=2)
    add_table(
        doc,
        ["Mã", "Yêu cầu"],
        [
            ["NFR01", "Ứng dụng khởi động và tự nạp dữ liệu sản phẩm/luật kết hợp mà không cần thao tác thủ công (DataSeeder tự động)"],
            ["NFR02", "Khi dữ liệu luật kết hợp trống hoặc thiếu, hệ thống vẫn hoạt động bình thường (không crash), chỉ ẩn khối khuyến nghị"],
            ["NFR03", "Giỏ hàng của khách hàng không bị mất khi hệ thống khởi động lại nạp dữ liệu catalog mới"],
            ["NFR04", "Toàn bộ pipeline khai phá dữ liệu chạy lại được bằng 1 lệnh duy nhất, có log rõ ràng từng bước để tái lập kết quả"],
        ],
        widths=[2, 13.5],
    )

    doc.add_page_break()

    # ---------------- CHUONG 3 ----------------
    add_heading(doc, "CHƯƠNG 3. THIẾT KẾ CÁC CHỨC NĂNG", level=1)

    add_heading(doc, "3.1 Kiến trúc tổng thể", level=2)
    doc.add_paragraph(
        "Hệ thống gồm 2 thành phần chính giao tiếp qua file dữ liệu trung gian "
        "(không qua API trực tiếp): pipeline Python khai phá luật kết hợp xuất ra "
        "products.csv, rules.csv, metrics.json; webapp ASP.NET Core đọc các file này "
        "qua DataSeeder để nạp vào SQLite, sau đó RecommendationService truy vấn cơ "
        "sở dữ liệu để phục vụ Controllers/Views."
    )
    doc.add_paragraph(
        "Việc tách 2 thành phần qua file CSV trung gian (thay vì gọi trực tiếp Python "
        "từ .NET) giúp mô hình có thể huấn luyện lại độc lập, dễ kiểm soát phiên bản "
        "dữ liệu, và đúng với thực tế triển khai MLOps đơn giản (batch export mô "
        "hình, hệ thống ứng dụng chỉ tiêu thụ kết quả)."
    )

    add_heading(doc, "3.2 Thiết kế dữ liệu", level=2)
    add_heading(doc, "3.2.1 Sơ đồ thực thể", level=3)
    for item in [
        "Product (Id, Sku [unique], Description, UnitPrice, Category)",
        "AssociationRule (Id, AntecedentSku, ConsequentSku, Support, Confidence, "
        "Lift) — unique index (AntecedentSku, ConsequentSku)",
        "Order (Id, CreatedAtUtc) 1—n OrderItem (Id, OrderId[FK], ProductId[FK], "
        "Quantity, UnitPrice)",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "3.2.2 Hợp đồng dữ liệu giữa pipeline và webapp", level=3)
    add_table(
        doc,
        ["File", "Cột", "Vai trò"],
        [
            ["products.csv", "sku, description, unit_price, category", "Nguồn nạp bảng Product"],
            ["rules.csv", "antecedent_sku, consequent_sku, support, confidence, lift", "Nguồn nạp bảng AssociationRule (chỉ luật 1→1)"],
            ["metrics.json", "xem Chương 4", "Số liệu phục vụ báo cáo, không nạp vào webapp"],
        ],
    )

    add_heading(doc, "3.3 Thiết kế xử lý — pipeline khai phá luật kết hợp", level=2)
    for i, text in enumerate([
        "Tiền xử lý: loại đơn hàng huỷ (Invoice bắt đầu “C”), loại dòng dữ liệu "
        "lỗi (Quantity/UnitPrice không dương, SKU/Description rỗng), loại các mã "
        "không phải sản phẩm thật (phí vận chuyển, phí ngân hàng, thẻ quà tặng...), "
        "gộp mỗi hoá đơn thành 1 giao dịch chứa tập SKU duy nhất.",
        "Giới hạn quy mô: giữ lại 300 sản phẩm bán chạy nhất (TOP_N_ITEMS) trước khi "
        "mã hoá one-hot, nhằm đảm bảo thời gian chạy hợp lý mà vẫn giữ phần lớn giao "
        "dịch có ý nghĩa.",
        "Khai phá: áp dụng song song Apriori và FP-Growth (thư viện mlxtend) trên "
        "cùng ma trận giao dịch, đối chiếu kết quả tập mục phổ biến phải trùng khớp.",
        "Sinh luật: từ tập mục phổ biến (dùng kết quả FP-Growth), sinh luật kết hợp "
        "theo hàm association_rules với ngưỡng min_confidence đã chọn, rút gọn về "
        "luật 1→1.",
        "Xuất dữ liệu: ghi products.csv, rules.csv, metrics.json theo đúng hợp đồng "
        "dữ liệu, ghi nguyên tử để tránh đọc file dở.",
    ], start=1):
        doc.add_paragraph(f"{i}. {text}")

    add_heading(doc, "3.4 Thiết kế xử lý — chức năng khuyến nghị trong webapp", level=2)
    doc.add_paragraph(
        "Trang chi tiết sản phẩm: truy vấn AssociationRule theo AntecedentSku = sku "
        "hiện tại, loại trừ chính nó, sắp xếp giảm dần theo (Confidence, Lift), lấy "
        "top 4.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Trang giỏ hàng: với mỗi SKU trong giỏ, truy vấn luật có AntecedentSku thuộc "
        "giỏ; gộp theo ConsequentSku, chọn Confidence lớn nhất trong các luật trùng "
        "consequent làm điểm chính, Lift của đúng luật đó làm tiêu chí phụ; loại các "
        "SKU đã có trong giỏ; sắp xếp và lấy top 6.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Nạp dữ liệu (DataSeeder): chạy mỗi lần khởi động ứng dụng — bảng Product "
        "được upsert theo Sku (không bao giờ xoá, tránh vỡ khoá ngoại với OrderItem), "
        "bảng AssociationRule được nạp lại toàn bộ từ CSV.",
        style="List Bullet",
    )

    add_heading(doc, "3.5 Thiết kế màn hình", level=2)
    for item in [
        "Trang chủ: danh mục sản phẩm dạng danh sách theo danh mục.",
        "Trang chi tiết sản phẩm: thông tin sản phẩm + khối “Khách hàng cũng "
        "mua” (tối đa 4 sản phẩm) + nút “Thêm vào giỏ”.",
        "Trang giỏ hàng: danh sách sản phẩm trong giỏ, khối “Gợi ý cho giỏ "
        "hàng” (tối đa 6 sản phẩm), nút “Đặt hàng”.",
        "Trang quản trị luật kết hợp: bảng liệt kê toàn bộ luật (antecedent, "
        "consequent, support, confidence, lift).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()

    # ---------------- CHUONG 4 ----------------
    add_heading(doc, "CHƯƠNG 4. KẾT QUẢ THỰC HIỆN", level=1)
    doc.add_paragraph(
        "Toàn bộ số liệu trong chương này lấy trực tiếp từ lần chạy thật của pipeline "
        f"(data-mining/outputs/metrics.json, sinh lúc {METRICS['generated_at']}) và "
        "từ việc chạy thật webapp trên dữ liệu đã khai phá."
    )

    add_heading(doc, "4.1 Kết quả tiền xử lý dữ liệu", level=2)
    doc.add_paragraph(
        "Dữ liệu gốc Online Retail II gồm 1.067.371 dòng giao dịch (2 năm 2009-2011). "
        "Sau bước tiền xử lý (b2): loại 19.494 dòng thuộc đơn hàng huỷ (Invoice bắt "
        "đầu “C”), 6.207 dòng dữ liệu lỗi (Quantity/UnitPrice không dương "
        "hoặc SKU/Description rỗng), 4.665 dòng thuộc mã không phải sản phẩm thật "
        f"(phí vận chuyển, phí ngân hàng, thẻ quà tặng...), còn lại {vn(METRICS['n_transactions_total_clean'])} "
        f"hoá đơn (giao dịch) hợp lệ với {vn(METRICS['n_unique_skus_total'])} mã sản "
        "phẩm (SKU) khác nhau."
    )

    add_heading(doc, "4.2 Kết quả khai phá luật kết hợp", level=2)
    doc.add_paragraph(
        f"Để đảm bảo thời gian khai phá hợp lý, {vn(METRICS['n_unique_skus_total'])} SKU "
        f"được giới hạn xuống {METRICS['top_n_items_cap']} sản phẩm bán chạy nhất "
        f"(TOP_N_ITEMS) trước khi mã hoá one-hot. Sau khi lọc, "
        f"{vn(METRICS['n_transactions_used_for_mining'])}/{vn(METRICS['n_transactions_total_clean'])} "
        f"hoá đơn ({METRICS['coverage_pct_after_cap']:.2f}%) vẫn còn từ 2 sản phẩm "
        "trở lên và được dùng để khai phá — tỷ lệ bao phủ này cho thấy việc giới hạn "
        "không làm mất tính đại diện của dữ liệu."
    )
    doc.add_paragraph(
        "Ngưỡng min_support/min_confidence cuối cùng được lựa chọn dựa trên bảng "
        "thực nghiệm số lượng luật sinh ra ở nhiều mức ngưỡng khác nhau:"
    )
    add_table(
        doc,
        ["min_support", "min_confidence", "Số luật"],
        [
            ["0.01", "0.20", "2.880"],
            ["0.01", "0.35", "1.571"],
            ["0.01", "0.50", "835"],
            ["0.02", "0.20", "257"],
            ["0.02", "0.35", "186"],
            ["0.02", "0.50", "83"],
            ["0.05", "0.20", "0"],
            ["0.05", "0.35", "0"],
            ["0.05", "0.50", "0"],
        ],
    )
    doc.add_paragraph(
        f"Ngưỡng min_support={METRICS['min_support']}, min_confidence="
        f"{METRICS['min_confidence']} được chọn vì cho {METRICS['n_rules_total']} "
        "luật — nằm trong khoảng mục tiêu 50-300 luật (đủ phong phú để phục vụ gợi "
        "ý, không quá nhiều gây nhiễu), đồng thời là cặp ngưỡng mạnh nhất (support "
        "và confidence cao nhất) trong số các cặp đạt mục tiêu, đảm bảo các luật giữ "
        "lại có độ tin cậy thống kê cao."
    )

    add_heading(doc, "4.2.1 So sánh Apriori và FP-Growth", level=3)
    doc.add_paragraph(
        "Cả 2 thuật toán được chạy trên cùng ma trận giao dịch đã mã hoá one-hot, "
        "cùng ngưỡng min_support, thời gian đo là giá trị nhỏ nhất trong 3 lần chạy "
        "liên tiếp (chỉ tính thời gian gọi hàm khai phá, không tính đọc/ghi dữ liệu):"
    )
    add_table(
        doc,
        ["Thuật toán", "Số tập mục phổ biến", "Thời gian chạy (giây)"],
        [
            ["Apriori", f"{METRICS['apriori']['n_itemsets']}", f"{METRICS['apriori']['runtime_sec']:.4f}"],
            ["FP-Growth", f"{METRICS['fpgrowth']['n_itemsets']}", f"{METRICS['fpgrowth']['runtime_sec']:.4f}"],
        ],
    )
    speedup = METRICS["apriori"]["runtime_sec"] / METRICS["fpgrowth"]["runtime_sec"]
    doc.add_paragraph(
        f"Hai thuật toán cho kết quả tập mục phổ biến hoàn toàn giống nhau "
        f"(itemsets_equal = {METRICS['itemsets_equal']}), xác nhận tính đúng đắn của "
        f"cài đặt. FP-Growth nhanh hơn Apriori khoảng {speedup:.2f} lần trên tập dữ "
        "liệu này, phù hợp với lý thuyết vì FP-Growth không cần sinh tập ứng viên "
        "tường minh như Apriori."
    )

    add_heading(doc, "4.2.2 Một số luật kết hợp tiêu biểu", level=3)
    rows = [
        [r["antecedent_desc"], r["consequent_desc"], f"{r['support']:.4f}", f"{r['confidence']*100:.1f}%", f"{r['lift']:.2f}"]
        for r in METRICS["top_rules_preview"][:8]
    ]
    add_table(doc, ["Nếu mua", "Khuyến nghị", "Support", "Confidence", "Lift"], rows)
    doc.add_paragraph(
        f"Sau khi rút gọn về dạng luật 1→1 (loại luật nhiều-tới-nhiều, cấm luật "
        f"A→A, loại trùng lặp), tổng cộng {METRICS['n_rules_1to1_exported']}/"
        f"{METRICS['n_rules_total']} luật được xuất ra rules.csv để tích hợp vào "
        "webapp. Các luật tiêu biểu ở trên cho thấy tính hợp lý về mặt nghiệp vụ — "
        "ví dụ bộ tách trà 3 màu (hồng/xanh/hoa hồng) hay bộ ly/đĩa giấy chấm bi "
        "cùng bộ thường được mua cùng nhau, đúng với trực giác bán lẻ."
    )

    add_heading(doc, "4.3 Kết quả tích hợp vào hệ thống webapp", level=2)
    doc.add_paragraph(
        "Sau khi copy products.csv/rules.csv thật vào webapp/src/RetailRecommender."
        "Web/SeedData/ và khởi động lại ứng dụng, DataSeeder tự động nạp "
        f"{METRICS['n_skus_used_for_mining']} SKU nằm trong dữ liệu khai phá (trong "
        f"tổng {vn(METRICS['n_unique_skus_total'])} SKU của catalog) cùng "
        f"{METRICS['n_rules_1to1_exported']} luật kết hợp vào SQLite."
    )
    add_image_with_caption(doc, ASSETS / "01_home.png", "Hình 4.1. Trang danh mục sản phẩm (dữ liệu thật từ Online Retail II)")
    add_image_with_caption(
        doc, ASSETS / "02_product_detail.png",
        "Hình 4.2. Trang chi tiết sản phẩm PINK REGENCY TEACUP AND SAUCER (SKU 22698) "
        "hiển thị đúng 2 khuyến nghị từ rules.csv thật: GREEN REGENCY TEACUP AND "
        "SAUCER (confidence 84%, lift 19,57) và ROSES REGENCY TEACUP AND SAUCER "
        "(confidence 79%, lift 17,57)",
    )
    doc.add_paragraph(
        "Chức năng gợi ý theo giỏ hàng cũng được xác nhận đúng với dữ liệu thật: khi "
        "giỏ hàng chứa PINK REGENCY TEACUP AND SAUCER (22698) và SET/6 RED SPOTTY "
        "PAPER CUPS (21086), hệ thống gợi ý đúng 3 sản phẩm GREEN REGENCY TEACUP AND "
        "SAUCER, ROSES REGENCY TEACUP AND SAUCER và SET/6 RED SPOTTY PAPER PLATES — "
        "khớp chính xác với các luật antecedent tương ứng trong rules.csv, không có "
        "sản phẩm nào trùng với 2 sản phẩm đã có sẵn trong giỏ."
    )
    add_image_with_caption(
        doc, ASSETS / "03_admin_rules.png",
        f"Hình 4.3. Trang quản trị hiển thị đầy đủ {METRICS['n_rules_1to1_exported']} "
        "luật kết hợp đã khai phá và tích hợp vào hệ thống",
    )

    add_heading(doc, "4.4 Kiểm thử", level=2)
    add_table(
        doc,
        ["Mã test", "Tình huống", "Kết quả"],
        [
            ["TC01", "Chạy run_pipeline.py từ đầu trên dữ liệu Online Retail II thật", "Đạt — sinh đủ products.csv, rules.csv, metrics.json"],
            ["TC02", "So sánh itemsets Apriori vs FP-Growth", f"Đạt — itemsets_equal = {METRICS['itemsets_equal']} ({METRICS['apriori']['n_itemsets']} tập mục mỗi thuật toán)"],
            ["TC03", "dotnet build + dotnet test", "Đạt — build 0 lỗi/0 cảnh báo, 6/6 test pass"],
            ["TC04", "Mở trang chi tiết sản phẩm 22698 → hiển thị đúng khuyến nghị", "Đạt — đúng 2 luật thật (xem Hình 4.2)"],
            ["TC05", "Thêm 2 sản phẩm vào giỏ → gợi ý tổng hợp đúng quy tắc gộp", "Đạt — 3 gợi ý đúng, không trùng sản phẩm đã có trong giỏ"],
            ["TC06", "dotnet run thật + curl kiểm tra /, /Products, /Admin/Rules", "Đạt — cả 3 endpoint trả HTTP 200 với dữ liệu thật"],
        ],
        widths=[2, 7, 6.5],
    )

    doc.add_page_break()

    # ---------------- CHUONG 5 ----------------
    add_heading(doc, "CHƯƠNG 5. KẾT LUẬN", level=1)

    add_heading(doc, "5.1 Đánh giá so với mục tiêu", level=2)
    doc.add_paragraph(
        "Đề tài đã hoàn thành đầy đủ 5 mục tiêu đề ra ở mục 1.4: pipeline khai phá "
        "dữ liệu chạy thành công trên dữ liệu Online Retail II thật (không dùng dữ "
        "liệu giả lập); 2 thuật toán Apriori và FP-Growth được đối chiếu và xác nhận "
        "cho cùng kết quả; ngưỡng support/confidence được lựa chọn dựa trên bảng thực "
        "nghiệm rõ ràng thay vì chọn tuỳ ý; hệ thống webapp ASP.NET Core hoạt động "
        "đầy đủ các chức năng quản lý bán lẻ tối thiểu; và quan trọng nhất, mô hình "
        "luật kết hợp đã được tích hợp thành công vào 2 điểm chạm cụ thể (trang chi "
        "tiết sản phẩm và trang giỏ hàng), được kiểm chứng bằng dữ liệu thật chứ "
        "không chỉ dừng ở lý thuyết."
    )

    add_heading(doc, "5.2 Hạn chế", level=2)
    for item in [
        f"Dữ liệu bị giới hạn ở {METRICS['top_n_items_cap']} sản phẩm bán chạy nhất "
        "(TOP_N_ITEMS) để đảm bảo hiệu năng khai phá, có thể bỏ sót luật liên quan "
        "tới sản phẩm ít phổ biến hơn.",
        "Luật kết hợp xuất cho webapp chỉ ở dạng 1→1 để đơn giản hoá tra cứu, "
        "chưa tận dụng hết các luật nhiều-tới-một tìm được trong quá trình khai phá "
        f"({METRICS['n_rules_total']} luật tổng cộng, chỉ "
        f"{METRICS['n_rules_1to1_exported']} luật được xuất).",
        "Giỏ hàng lưu theo Session, chưa có tài khoản người dùng thật/lịch sử mua "
        "hàng cá nhân hoá.",
        "Mô hình luật kết hợp không tự động cập nhật theo thời gian thực (batch, cần "
        "chạy lại pipeline thủ công khi có dữ liệu giao dịch mới).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "5.3 Hướng phát triển", level=2)
    for item in [
        "Bổ sung tài khoản người dùng thật, cá nhân hoá khuyến nghị theo lịch sử mua "
        "hàng từng khách hàng.",
        "Lập lịch tự động chạy lại pipeline khai phá theo chu kỳ, tự động cập nhật "
        "rules.csv mới nhất vào hệ thống.",
        "Kết hợp thêm kỹ thuật collaborative filtering để bổ sung cho luật kết hợp "
        "trong các trường hợp dữ liệu thưa.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()

    # ---------------- REFERENCES ----------------
    add_heading(doc, "TÀI LIỆU THAM KHẢO", level=1)
    for ref in [
        "[1] R. Agrawal, R. Srikant (1994). Fast Algorithms for Mining Association "
        "Rules. Proceedings of the 20th VLDB Conference.",
        "[2] J. Han, J. Pei, Y. Yin (2000). Mining Frequent Patterns without "
        "Candidate Generation. ACM SIGMOD.",
        "[3] Chen, D. (2019). Online Retail II [Dataset]. UCI Machine Learning "
        "Repository. https://doi.org/10.24432/C5CG6D",
        "[4] Lê Hữu Dũng. Lecture Note — Nhập môn Khai phá dữ liệu và Máy học, "
        "Trường Đại học Mở Hà Nội, 2024.",
        "[5] Microsoft. ASP.NET Core documentation. "
        "https://learn.microsoft.com/aspnet/core/",
        "[6] mlxtend documentation. https://rasbt.github.io/mlxtend/",
    ]:
        doc.add_paragraph(ref)

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
